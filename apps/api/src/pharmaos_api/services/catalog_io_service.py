"""Catalog import/export with batch quantities (super_admin only).

One file format, three entry points (CLAUDE.md "قالب ملف استيراد/تصدير الكاتلوج"):

- Export   (GET  /catalog/export)         → the catalog as .xlsx, optionally with
                                            one row per ACTIVE batch of a chosen
                                            branch (the quantities).
- Template (GET  /catalog/import-template)→ the empty header sheet + an Arabic
                                            instructions sheet — generated from
                                            the same _COLUMNS constant the parser
                                            validates against, so they can never
                                            drift apart.
- Import   (POST /catalog/import)         → multipart rows (xlsx or csv parsed by
                                            the router) with a dry_run preview.

Matching: barcode/gtin exact → exact (case-insensitive) name → CREATE. Never
fuzzy; an ambiguous key is rejected, not guessed. Quantities NEVER live on the
medication (forbidden action #15): a row with quantity > 0 creates a batch on
the selected branch through inventory_service.receive_batch (batch row +
purchase_in movement + branch_inventory cache delta in the caller's tx).

Round-trip contract: an exported file re-imports cleanly — duplicate detection
is per (medication, batch_number) pair, NOT per medication, because the export
legitimately repeats the medication across its batch rows. Updates are
idempotent: empty cells keep the existing value, trade_name is a match KEY
(never renamed by import), and a repeated medication in the file is harmless.

Rejected rows always carry a named reason with their spreadsheet row number —
nothing is silently dropped, and a rejected row applies nothing.
"""

import datetime as dt
import io
import uuid
from collections.abc import Sequence
from decimal import Decimal, InvalidOperation
from typing import Any

from sqlalchemy import and_, insert, select, text
from sqlalchemy.ext.asyncio import AsyncSession

from pharmaos_api.audit import AuditAction
from pharmaos_api.errors import ApiError, ErrorCode
from pharmaos_api.models import (
    Branch,
    Medication,
    MedicationBarcode,
    MedicationBatch,
    MedicationPackaging,
    MedicationPriceHistory,
    User,
)
from pharmaos_api.services import audit_service, inventory_service, seed_service

# The single source of truth for the file format — the template, the export and
# the import parser all read this tuple (CLAUDE.md documents the same list).
_COLUMNS = (
    "trade_name_en",
    "trade_name_ar",
    "scientific_name",
    "manufacturer",
    "drug_class",
    "route",
    "gtin",
    "price_egp",
    "branch",
    "batch_number",
    "expiry_date",
    "quantity",
    "purchase_price",
)
# Columns of an existing medication a row may refresh (trade_name is the match
# key and is deliberately absent — import never renames).
_UPDATABLE_FIELDS = (
    "trade_name_ar",
    "scientific_name",
    "manufacturer",
    "drug_class",
    "route",
    "gtin",
)

CATALOG_IMPORT_MAX_BYTES = 10 * 1024 * 1024
CATALOG_IMPORT_MAX_ROWS = 30_000
_BATCH = 1000
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_PRICE_2 = Decimal("0.01")
_QTY_3 = Decimal("0.001")


# ------------------------------------------------------------------ export --


async def export_catalog_xlsx(
    session: AsyncSession, *, branch_id: uuid.UUID | None = None
) -> tuple[bytes, int]:
    """The full catalog as an .xlsx workbook.

    Without a branch: one row per medication, batch columns empty (the pure
    catalog template). With a branch: medications holding ACTIVE batches in
    that branch get one row per batch (the quantities); the rest still get a
    single catalog row, so the file is always a complete round-trip template.
    Returns (bytes, data_row_count).
    """
    branch: Branch | None = None
    if branch_id is not None:
        branch = await session.get(Branch, branch_id)
        if branch is None or branch.is_deleted or not branch.is_active:
            raise ApiError(ErrorCode.VALIDATION_FAILED, 422, message="Unknown branch.")

    med_rows = (
        await session.execute(
            select(
                Medication.id,
                Medication.trade_name,
                Medication.trade_name_ar,
                Medication.scientific_name,
                Medication.manufacturer,
                Medication.drug_class,
                Medication.route,
                Medication.gtin,
                MedicationPackaging.selling_price,
            )
            .outerjoin(
                MedicationPackaging,
                and_(
                    MedicationPackaging.medication_id == Medication.id,
                    MedicationPackaging.level == 1,
                    MedicationPackaging.is_deleted.is_(False),
                ),
            )
            .where(Medication.is_deleted.is_(False), Medication.is_active.is_(True))
            .order_by(Medication.trade_name)
        )
    ).all()

    batches_by_med: dict[uuid.UUID, list[Any]] = {}
    if branch is not None:
        for batch_row in (
            await session.execute(
                select(
                    MedicationBatch.medication_id,
                    MedicationBatch.batch_number,
                    MedicationBatch.expiry_date,
                    MedicationBatch.quantity,
                    MedicationBatch.purchase_price,
                )
                .where(
                    MedicationBatch.branch_id == branch_id,
                    MedicationBatch.status == "active",
                )
                .order_by(MedicationBatch.expiry_date)
            )
        ).all():
            batches_by_med.setdefault(batch_row.medication_id, []).append(batch_row)

    import openpyxl
    from openpyxl.cell import WriteOnlyCell

    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet("القالب")
    ws.append(list(_COLUMNS))

    def _text_row(values: list[Any]) -> None:
        """append() would store '='-prefixed text as FORMULAS (openpyxl value
        inference), which read back as None under data_only — breaking the
        round-trip for names like the security corpus's '=HYPERLINK(...)'.
        Every export cell is forced to a plain string cell."""
        cells = []
        for value in values:
            cell = WriteOnlyCell(ws, value="" if value is None else value)
            cell.data_type = "s"
            cells.append(cell)
        ws.append(cells)

    def _catalog_cells(row: Any, branch_name: str) -> list[Any]:
        return [
            row.trade_name,
            row.trade_name_ar or "",
            row.scientific_name or "",
            row.manufacturer or "",
            row.drug_class or "",
            row.route or "",
            row.gtin or "",
            str(row.selling_price) if row.selling_price is not None else "",
            branch_name,
            "",
            "",
            "",
            "",
        ]

    branch_name = branch.name if branch is not None else ""
    count = 0
    for med_row in med_rows:
        batch_rows = batches_by_med.get(med_row.id) if branch is not None else None
        if not batch_rows:
            _text_row(_catalog_cells(med_row, branch_name))
            count += 1
            continue
        for b in batch_rows:
            cells = _catalog_cells(med_row, branch_name)
            cells[_COLUMNS.index("batch_number")] = b.batch_number
            cells[_COLUMNS.index("expiry_date")] = b.expiry_date.isoformat()
            cells[_COLUMNS.index("quantity")] = str(b.quantity)
            cells[_COLUMNS.index("purchase_price")] = str(b.purchase_price)
            _text_row(cells)
            count += 1

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), count


# ---------------------------------------------------------------- template --


def build_import_template_xlsx() -> bytes:
    """The empty import template: a header-only sheet (nothing importable by
    accident) + an Arabic instructions sheet describing every column and rule."""
    import openpyxl

    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet("القالب")
    ws.append(list(_COLUMNS))

    doc = wb.create_sheet("تعليمات")
    doc.append(["تعليمات استيراد كاتلوج الأدوية — PharmaOS"])
    doc.append([])
    doc.append(
        [
            "1",
            "اكتب كل دواء في صف واحد ولا تغيّر أسماء الأعمدة الإنجليزية"
            " في الصف الأول ولا تحذفها.",
        ]
    )
    doc.append(
        [
            "2",
            "trade_name_en إلزامي وهو مفتاح المطابقة — الاستيراد لا يعيد تسمية دواء قائم؛"
            " عدّل الاسم يدويًا من محرر الكاتلوج.",
        ]
    )
    doc.append(
        [
            "3",
            "المطابقة بترتيب: gtin/باركود ثم الاسم الإنجليزي/العربي مطابقًا تمامًا"
            " (غير حساس لحالة الأحرف). دواء غير مطابق ⇒ يُنشأ جديدًا."
            " اسم يطابق أكثر من دواء ⇒ الصف مرفوض (ambiguous_name).",
        ]
    )
    doc.append(
        [
            "4",
            "لإضافة مخزون: املأ quantity (أكبر من صفر) مع batch_number و expiry_date"
            " بصيغة YYYY-MM-DD — تُنشأ دفعة جديدة على الفرع الذي تختاره"
            " في نافذة الاستيراد وتُقيَّد في سجل الحركات.",
        ]
    )
    doc.append(["5", "اترك quantity فارغًا لتحديث بيانات الكاتلوج فقط دون أي حركة مخزون."])
    doc.append(
        [
            "6",
            "عمود branch اختياري: اتركه فارغًا أو اكتب اسم الفرع نفسه المختار"
            " في نافذة الاستيراد وإلا رُفض الصف (branch_mismatch).",
        ]
    )
    doc.append(
        [
            "7",
            "price_egp يحدّث سعر العلبة (المستوى 1) لدواء قائم ويُسجَّل في سجل تاريخ"
            " الأسعار. الخلية الفارغة تعني: اترك القيمة الحالية كما هي.",
        ]
    )
    doc.append(
        [
            "8",
            "purchase_price اختياري ويعود للدفعة (افتراضه 0)."
            " تاريخ منقضٍ أو كمية غير صالحة ⇒ الصف مرفوض.",
        ]
    )
    doc.append(["9", "الحدود: 30,000 صف و10MB. الملف المصدَّر من «تصدير» جاهز للاستيراد مباشرة."])
    doc.append(["10", "استخدم «معاينة» أولًا — لا يُكتب شيء في القاعدة قبل «تنفيذ الاستيراد»."])
    doc.append([])
    doc.append(["العمود", "إلزامي؟", "الوصف ومثال"])
    descriptions = {
        "trade_name_en": "الاسم التجاري الإنجليزي — مفتاح المطابقة (مثال: Panadol Extra)",
        "trade_name_ar": "الاسم التجاري العربي (مثال: بانادول إكسترا)",
        "scientific_name": "المادة الفعالة (مثال: Paracetamol)",
        "manufacturer": "الشركة المصنعة (مثال: GSK)",
        "drug_class": "التصنيف الدوائي (مثال: Analgesic)",
        "route": "طريقة الإعطاء (مثال: Oral)",
        "gtin": "باركود GTIN حتى 14 رقمًا (مثال: 6221031420016)",
        "price_egp": "سعر بيع العلبة (مثال: 25.50)",
        "branch": "اسم الفرع كما في النظام — اختياري (مثال: الفرع الرئيسي)",
        "batch_number": "رقم التشغيلة — إلزامي مع الكمية (مثال: B2401)",
        "expiry_date": "تاريخ الانتهاء YYYY-MM-DD — إلزامي مع الكمية ومستقبلي (مثال: 2027-05-31)",
        "quantity": "الكمية — أكبر من صفر (مثال: 100)",
        "purchase_price": "سعر شراء الدفعة — اختياري (مثال: 20.00)",
    }
    for col in _COLUMNS:
        doc.append([col, "إلزامي" if col == "trade_name_en" else "اختياري", descriptions[col]])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ import --


def _parse_price(raw: str) -> Decimal | None:
    try:
        price = Decimal(raw).quantize(_PRICE_2)
    except InvalidOperation:
        return None
    return price if price >= 0 else None


def _parse_date(raw: str) -> dt.date | None:
    """ISO strings (the router normalizes Excel datetime cells to ISO dates)."""
    try:
        return dt.date.fromisoformat(raw)
    except ValueError:
        return None


def _load_resolution_maps_sync(
    rows: Sequence[Any],
    barcode_rows: Sequence[Any],
    l1_rows: Sequence[Any],
) -> tuple[
    dict[str, list[uuid.UUID]],
    dict[str, list[uuid.UUID]],
    dict[str, uuid.UUID],
    dict[uuid.UUID, tuple[str, str | None]],
    dict[uuid.UUID, tuple[Any, Any, Decimal]],
]:
    """Build the bulk matching maps from pre-fetched rows — one pass over the
    catalog instead of per-row queries (a file may carry ~25k rows). Names are
    casefolded exact keys taken from BOTH trade_name and trade_name_ar."""
    name_map: dict[str, list[uuid.UUID]] = {}
    gtin_map: dict[str, list[uuid.UUID]] = {}
    names: dict[uuid.UUID, tuple[str, str | None]] = {}
    for med_id, trade_name, trade_name_ar, gtin in rows:
        names[med_id] = (trade_name, trade_name_ar)
        for key in {trade_name.casefold(), (trade_name_ar or "").casefold()}:
            if key:
                name_map.setdefault(key, []).append(med_id)
        if gtin:
            gtin_map.setdefault(gtin, []).append(med_id)
    barcode_map: dict[str, uuid.UUID] = {r[0]: r[1] for r in barcode_rows}
    l1_packaging: dict[uuid.UUID, tuple[Any, Any, Decimal]] = {r[0]: r for r in l1_rows}
    return name_map, gtin_map, barcode_map, names, l1_packaging


def _resolve(
    name_map: dict[str, list[uuid.UUID]],
    gtin_map: dict[str, list[uuid.UUID]],
    barcode_map: dict[str, uuid.UUID],
    *,
    name_key: str,
    gtin: str,
) -> tuple[uuid.UUID | None, str | None]:
    """Barcode/gtin exact → exact name → None (create). Returns
    (medication_id, None) on a unique hit or (None, 'ambiguous_name')."""
    if gtin:
        hit = barcode_map.get(gtin)
        if hit is not None:
            return hit, None
        gtin_hits = gtin_map.get(gtin, [])
        if len(gtin_hits) == 1:
            return gtin_hits[0], None
        if len(gtin_hits) > 1:
            return None, "ambiguous_name"
    name_hits = name_map.get(name_key, [])
    if len(name_hits) == 1:
        return name_hits[0], None
    if len(name_hits) > 1:
        return None, "ambiguous_name"
    return None, None


async def import_catalog(
    session: AsyncSession,
    *,
    actor: User,
    branch_id: uuid.UUID,
    rows: list[dict[str, str]],
    dry_run: bool = True,
    filename: str | None = None,
) -> dict[str, object]:
    """Upsert catalog rows + create batches on the selected branch.

    dry_run validates everything and writes nothing; execute re-validates
    fully (a preview may be stale) then applies: bulk medication/packaging
    inserts, idempotent field/price updates, and one receive_batch per
    quantity row — committed in chunks so a 25k-row file stays one clean
    sequence of transactions.
    """
    branch = await session.get(Branch, branch_id)
    if branch is None or branch.is_deleted or not branch.is_active:
        raise ApiError(ErrorCode.VALIDATION_FAILED, 422, message="Unknown branch.")
    branch_key = branch.name.casefold()
    if len(rows) > CATALOG_IMPORT_MAX_ROWS:
        raise ApiError(ErrorCode.VALIDATION_FAILED, 422, message="Too many rows.")

    med_rows = (
        await session.execute(
            select(
                Medication.id,
                Medication.trade_name,
                Medication.trade_name_ar,
                Medication.gtin,
            ).where(Medication.is_deleted.is_(False))
        )
    ).all()
    barcode_rows = (
        await session.execute(
            select(MedicationBarcode.barcode, MedicationBarcode.medication_id).where(
                MedicationBarcode.is_deleted.is_(False)
            )
        )
    ).all()
    l1_rows = (
        await session.execute(
            select(
                MedicationPackaging.medication_id,
                MedicationPackaging.id,
                MedicationPackaging.selling_price,
            ).where(
                MedicationPackaging.level == 1,
                MedicationPackaging.is_deleted.is_(False),
            )
        )
    ).all()
    name_map, gtin_map, barcode_map, names, l1_packaging = _load_resolution_maps_sync(
        med_rows, barcode_rows, l1_rows
    )

    today = dt.date.today()
    parsed: list[dict[str, Any]] = []
    seen_batches: set[tuple[Any, str]] = set()
    for i, raw in enumerate(rows):
        excel_row = i + 2  # 1-based + header row
        entry: dict[str, Any] = {"row": excel_row, "ok": False, "name": None}

        name = (raw.get("trade_name_en") or "").strip()
        if not name:
            entry["reason"] = "empty_name"
            parsed.append(entry)
            continue
        entry["name"] = name
        name_key = name.casefold()
        name_ar = (raw.get("trade_name_ar") or "").strip()
        gtin = (raw.get("gtin") or "").strip()

        file_branch = (raw.get("branch") or "").strip()
        if file_branch and file_branch.casefold() != branch_key:
            entry["reason"] = "branch_mismatch"
            parsed.append(entry)
            continue

        price: Decimal | None = None
        price_raw = (raw.get("price_egp") or "").strip()
        if price_raw:
            price = _parse_price(price_raw)
            if price is None:
                entry["reason"] = "bad_price"
                parsed.append(entry)
                continue

        med_id, err = _resolve(name_map, gtin_map, barcode_map, name_key=name_key, gtin=gtin)
        if err is not None:
            entry["reason"] = err
            parsed.append(entry)
            continue
        identity: Any = med_id if med_id is not None else f"new:{name_key}"

        batch: dict[str, Any] | None = None
        qty_raw = (raw.get("quantity") or "").strip()
        if qty_raw:
            try:
                quantity = Decimal(qty_raw).quantize(_QTY_3)
            except InvalidOperation:
                quantity = Decimal("-1")
            if quantity <= 0:
                entry["reason"] = "bad_quantity"
                parsed.append(entry)
                continue
            batch_number = (raw.get("batch_number") or "").strip()
            if not batch_number:
                entry["reason"] = "missing_batch_number"
                parsed.append(entry)
                continue
            expiry_raw = (raw.get("expiry_date") or "").strip()
            expiry = _parse_date(expiry_raw) if expiry_raw else None
            if expiry is None:
                entry["reason"] = "missing_expiry" if not expiry_raw else "bad_expiry"
                parsed.append(entry)
                continue
            if expiry <= today:
                entry["reason"] = "expired_batch"
                parsed.append(entry)
                continue
            purchase_price = Decimal("0.00")
            pp_raw = (raw.get("purchase_price") or "").strip()
            if pp_raw:
                parsed_pp = _parse_price(pp_raw)
                if parsed_pp is None:
                    entry["reason"] = "bad_purchase_price"
                    parsed.append(entry)
                    continue
                purchase_price = parsed_pp
            batch_key = (identity, batch_number)
            if batch_key in seen_batches:
                entry["reason"] = "duplicate_batch_in_file"
                parsed.append(entry)
                continue
            seen_batches.add(batch_key)
            batch = {
                "quantity": quantity,
                "batch_number": batch_number[:50],
                "expiry_date": expiry,
                "purchase_price": purchase_price,
            }

        fields = {
            "trade_name_ar": name_ar[:255] or None,
            "scientific_name": (raw.get("scientific_name") or "").strip()[:255] or None,
            "manufacturer": (raw.get("manufacturer") or "").strip()[:255] or None,
            "drug_class": (raw.get("drug_class") or "").strip()[:100] or None,
            "route": (raw.get("route") or "").strip()[:50] or None,
            "gtin": gtin[:14] or None,
        }
        entry["ok"] = True
        entry["med_id"] = med_id
        entry["identity"] = identity
        entry["is_create"] = med_id is None
        entry["name_key"] = name_key
        entry["fields"] = fields
        entry["price"] = price
        entry["batch"] = batch
        # A pure name-only row resolves but changes nothing — don't count it
        # as an update (the trade_name key itself is never written).
        entry["touches"] = (
            any(fields[f] for f in _UPDATABLE_FIELDS) or price is not None or batch is not None
        )
        parsed.append(entry)

    rejected = [
        {"row": p["row"], "trade_name": p.get("name"), "reason": p["reason"]}
        for p in parsed
        if not p["ok"]
    ]
    valid = [p for p in parsed if p["ok"]]

    if dry_run:
        to_create = {p["identity"] for p in valid if p["is_create"]}
        to_update = {p["med_id"] for p in valid if not p["is_create"] and p["touches"]}
        batches = [p for p in valid if p["batch"] is not None]
        preview = [
            {
                "row": p["row"],
                "action": ("create" if p["is_create"] else "update")
                + ("+batch" if p["batch"] is not None else ""),
                "trade_name": names[p["med_id"]][0] if p["med_id"] is not None else p["name"],
                "quantity": str(p["batch"]["quantity"]) if p["batch"] else None,
            }
            for p in valid[:50]
        ]
        return {
            "branch_id": str(branch_id),
            "dry_run": True,
            "total": len(rows),
            "to_create": len(to_create),
            "to_update": len(to_update),
            "batches": len(batches),
            "quantity_total": str(sum((p["batch"]["quantity"] for p in batches), Decimal("0"))),
            "valid": preview,
            "rejected": rejected,
        }

    # --------------- execute (fully re-validated above in THIS call) --------
    # Releases the unit upsert row lock immediately (same as the seeder); safe
    # here because nothing has been written yet in this transaction.
    unit_id = await seed_service._box_unit_id(session)  # noqa: SLF001 — house style (svc._UNSET)

    # 1) Creates — one medication per new name key (later rows with the same
    #    key are updates of the row that created it; no phantom duplicates).
    unique_creates: dict[str, dict[str, Any]] = {}
    for p in valid:
        if p["is_create"]:
            unique_creates.setdefault(p["identity"], p)
    created_med_ids: dict[str, uuid.UUID] = {}
    create_list = list(unique_creates.values())
    for chunk_start in range(0, len(create_list), _BATCH):
        chunk = create_list[chunk_start : chunk_start + _BATCH]
        med_inserted = (
            await session.execute(
                insert(Medication).returning(Medication.id, Medication.trade_name),
                [
                    {
                        "trade_name": p["name"][:255],
                        **p["fields"],
                        "created_by": actor.id,
                        "updated_by": actor.id,
                    }
                    for p in chunk
                ],
            )
        ).all()
        by_name = {name.casefold(): med_id for med_id, name in med_inserted}
        pkg_ids = (
            await session.execute(
                insert(MedicationPackaging).returning(
                    MedicationPackaging.id, MedicationPackaging.medication_id
                ),
                [
                    {
                        "medication_id": by_name[p["name_key"]],
                        "level": 1,
                        "unit_id": unit_id,
                        "name_ar": seed_service._BOX_UNIT_NAME,  # noqa: SLF001
                        "qty_in_parent": None,
                        "is_sellable": True,
                        "selling_price": p["price"] if p["price"] is not None else Decimal("0.00"),
                        "is_default_sale": True,
                        "price_source": "import",
                    }
                    for p in chunk
                ],
            )
        ).all()
        pkg_by_med = {med_id: pkg_id for pkg_id, med_id in pkg_ids}
        await session.execute(
            insert(MedicationPriceHistory),
            [
                {
                    "medication_id": med_id,
                    "packaging_id": pkg_by_med[med_id],
                    "old_price": None,
                    "new_price": p["price"] if p["price"] is not None else Decimal("0.00"),
                    "price_source": "import",
                }
                for pkg_id, med_id in pkg_ids
            ],
        )
        for p in chunk:
            created_med_ids[p["identity"]] = by_name[p["name_key"]]
        await session.commit()

    # 2) Field updates — merged first-wins per medication, empty cells keep
    #    the current value (COALESCE); trade_name is never touched.
    field_updates: dict[uuid.UUID, dict[str, Any]] = {}
    price_updates: dict[uuid.UUID, Decimal] = {}
    for p in valid:
        if p["is_create"] or not p["touches"]:
            continue
        med_id = p["med_id"]
        if p["price"] is not None and med_id in l1_packaging:
            price_updates.setdefault(med_id, p["price"])
        slot = field_updates.setdefault(med_id, dict.fromkeys(_UPDATABLE_FIELDS, None))
        for f in _UPDATABLE_FIELDS:
            if slot[f] is None and p["fields"][f]:
                slot[f] = p["fields"][f]
    field_updates = {mid: f for mid, f in field_updates.items() if any(f.values())}

    field_list = list(field_updates.items())
    for chunk_start in range(0, len(field_list), _BATCH):
        field_chunk = field_list[chunk_start : chunk_start + _BATCH]
        await session.execute(
            text("""
                UPDATE medications SET
                    trade_name_ar   = COALESCE(:trade_name_ar, trade_name_ar),
                    scientific_name = COALESCE(:scientific_name, scientific_name),
                    manufacturer    = COALESCE(:manufacturer, manufacturer),
                    drug_class      = COALESCE(:drug_class, drug_class),
                    route           = COALESCE(:route, route),
                    gtin            = COALESCE(:gtin, gtin),
                    updated_by      = :actor
                WHERE id = :id
                """),
            [{"id": med_id, "actor": actor.id, **vals} for med_id, vals in field_chunk],
        )
        await session.commit()

    price_list = list(price_updates.items())
    for chunk_start in range(0, len(price_list), _BATCH):
        price_chunk = price_list[chunk_start : chunk_start + _BATCH]
        await session.execute(
            text("UPDATE medication_packaging SET selling_price = :price WHERE id = :pkg_id"),
            [{"pkg_id": l1_packaging[med_id][1], "price": price} for med_id, price in price_chunk],
        )
        await session.execute(
            insert(MedicationPriceHistory),
            [
                {
                    "medication_id": med_id,
                    "packaging_id": l1_packaging[med_id][1],
                    "old_price": l1_packaging[med_id][2],
                    "new_price": price,
                    "price_source": "import",
                }
                for med_id, price in price_chunk
            ],
        )
        await session.commit()

    # 3) Batches — one receive_batch per quantity row (batch row + purchase_in
    #    movement + cache delta inside this tx; committed in chunks).
    applied_batches = 0
    quantity_total = Decimal("0")
    apply_errors = list(rejected)
    applied_rows: list[dict[str, object]] = []
    for p in valid:
        if p["batch"] is None:
            continue
        med_id = created_med_ids.get(p["identity"], p["med_id"])
        try:
            await inventory_service.receive_batch(
                session,
                actor=actor,
                branch_id=branch_id,
                medication_id=med_id,
                batch_number=p["batch"]["batch_number"],
                expiry_date=p["batch"]["expiry_date"],
                quantity=p["batch"]["quantity"],
                purchase_price=p["batch"]["purchase_price"],
                reference_type="catalog_import",
            )
        except ApiError as e:
            apply_errors.append(
                {"row": p["row"], "trade_name": p["name"], "reason": "apply_failed", "code": e.code}
            )
            continue
        applied_batches += 1
        quantity_total += p["batch"]["quantity"]
        applied_rows.append({"row": p["row"], "trade_name": p["name"]})
        if applied_batches % _BATCH == 0:
            await session.commit()

    await audit_service.record(
        session,
        AuditAction.CATALOG_IMPORTED,
        actor=actor,
        branch_id=branch_id,
        entity_type="catalog",
        metadata={
            "created": len(created_med_ids),
            "updated": len(field_updates) + len(price_updates),
            "batches": applied_batches,
            "quantity_total": str(quantity_total),
            "filename": filename,
        },
    )
    await session.commit()

    return {
        "branch_id": str(branch_id),
        "dry_run": False,
        "total": len(rows),
        "created": len(created_med_ids),
        "updated": len(field_updates) + len(price_updates),
        "batches": applied_batches,
        "quantity_total": str(quantity_total),
        "rows": applied_rows[:50],
        "rejected": sorted(apply_errors, key=lambda e: e["row"]),
    }
